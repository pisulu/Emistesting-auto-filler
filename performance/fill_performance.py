#!/usr/bin/env python3
"""Guided term-performance (marks) filler for the SEMIS lifecycle test.

Run it from the folder that holds your data (tester workbooks, calendar, templates),
or point to them with --data-dir / --templates.

For each standard (and academic year) in the tester's test data the script:
  1. tells you what to download from SEMIS (standard, year, start and closing date from the calendar),
  2. fills the downloaded "SEMIS - Learners Performance - ..." template with the term scores from
     the Marks sheet of the tester workbook,
  3. writes a report with the scores entered, the grade level and term remark SEMIS should show,
     and anything that could not be entered,
  4. asks whether to continue with the next standard.

Usage:
    python fill_performance.py                   # interactive: asks which tester(s), then guides each standard
    python fill_performance.py --tester 2        # Tester 2 only (SEMIS_Test_Tester_2.xlsx)
    python fill_performance.py --tester all      # all testers into the same templates
    python fill_performance.py --templates "C:\\Users\\me\\Downloads"
    python fill_performance.py --subject-map my_subjects.json
    python fill_performance.py --yes             # never ask, skip standards whose template is missing

Rules applied:
  * Only the score cells are written. "Subject Grade" (Level 1-4) and "Term Remarks" are formulas in the
    template and fill themselves in; the report shows what they should come to.
  * The template decides which subjects exist (they differ for infants, junior and senior classes).
    Test-data subjects are mapped onto SEMIS subjects (see SUBJECT_MAP); a test-data subject with no
    column in the template is not entered and is listed in the report.
  * The performance template does not say which academic year it is for. The year is taken from the
    file name when it contains one (e.g. "... - 2019.xlsx"), otherwise from the learners in the file.
  * The tester workbook itself is never modified.
"""
import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))   # shared helpers live in the repo root

import openpyxl
from openpyxl.styles import Font

from semis_common import (SKIP_SHEETS, add_common_args, ask, code_of, dmy, prepare, sheet, std_num)

# SEMIS subject -> test-data subjects to take the score from (first one that has a score wins).
# Edit here or pass --subject-map <json file> with the same shape.
SUBJECT_MAP = {
    "Mathematics": ["Mathematics"],
    "Chichewa": ["Chichewa"],
    "English": ["English"],
    "Expressive Arts": ["Expressive Arts"],
    "Agriculture": ["Agriculture"],
    "Primary Science": ["Science & Technology", "Primary Science"],
    "Social Studies": ["Social & Environmental Sciences", "Social Studies"],
    "Bible Knowledge": ["Religious Education", "Life Skills"],
}
# Test-data remark -> the SEMIS remark it is closest to (used only for the comparison column).
REMARK_EQUIV = {"excellent": "EXCELLENT", "very good": "GOOD", "good": "GOOD", "pass": "AVERAGE",
                "needs improvement": "NEEDS SUPPORT"}


def norm(name):
    return re.sub(r"[^a-z]", "", str(name).lower())


def level_of(score):
    """Mirrors the template's 'Subject Grade' formula."""
    return "Level 1" if score < 40 else "Level 2" if score < 60 else "Level 3" if score < 80 else "Level 4"


def remark_of(avg):
    """Mirrors the template's 'Term Remarks' formula."""
    return "EXCELLENT" if avg >= 80 else "GOOD" if avg >= 60 else "AVERAGE" if avg >= 40 else "NEEDS SUPPORT"


# --------------------------------------------------------------------------- test data
@dataclass
class Record:
    tester: int
    row: int
    id: str
    name: str
    label: str
    term: int
    standard: str
    school: str
    scores: dict                 # test-data subject -> score
    average: float
    remark: str
    problems: list = field(default_factory=list)


def load_records(path, tester):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["Marks"]
    rows = list(ws.iter_rows(min_row=4, values_only=True))
    wb.close()
    head = [str(h or "").strip() for h in rows[0]]
    first_grade = next(i for i, h in enumerate(head) if h.lower().endswith(" grade"))
    subjects = {i: head[i] for i in range(6, first_grade)}
    avg_i, remark_i = head.index("Average"), head.index("Remark")
    out = []
    for n, r in enumerate(rows[1:], start=5):
        if not r or not r[0]:
            continue
        rec = Record(tester=tester, row=n, id=str(r[0]), name=r[1], label=r[2], term=int(r[3]),
                     standard=r[4], school=r[5], scores={}, average=r[avg_i], remark=r[remark_i])
        for i, subject in subjects.items():
            v = r[i]
            if v is None or str(v).strip() == "":
                continue
            try:
                score = float(v)
            except ValueError:
                rec.problems.append(f"{subject}: '{v}' is not a number (not entered)")
                continue
            if not 0 <= score <= 100:
                rec.problems.append(f"{subject}: {v} is outside 0-100 (entered as is)")
            if score != int(score):
                rec.problems.append(f"{subject}: {v} is not a whole number (rounded)")
            rec.scores[subject] = int(round(score))
        out.append(rec)
    return out


@dataclass
class Job:
    std: int
    standard: str
    label: str
    code: str
    records: list

    @property
    def learners(self):
        return sorted({r.id for r in self.records})


def build_jobs(records):
    groups = defaultdict(list)
    for r in records:
        groups[(std_num(r.standard), code_of(r.label), r.standard, r.label)].append(r)
    return [Job(k[0], k[2], k[3], k[1], v) for k, v in sorted(groups.items())]


# --------------------------------------------------------------------------- templates
@dataclass
class Template:
    path: Path
    standard: str
    stream: str
    ids: set
    hint: str            # academic year code taken from the file name, or ""


def year_hint(name):
    m = re.search(r"(20\d\d)\s*-\s*(20\d\d)", name)
    if m:
        return m.group(2)
    m = re.search(r"(?<!\d)(20\d\d)(?!\d)", name)
    return m.group(1) if m else ""


def data_sheet(wb):
    return next(ws for ws in wb.worksheets if ws.title not in SKIP_SHEETS)


def peek_template(path):
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = data_sheet(wb)
        rows = [r for r in ws.iter_rows(min_row=4, max_col=4, values_only=True) if r and r[3]]
        wb.close()
        return Template(Path(path), str(rows[0][1]), str(rows[0][2]),
                        {str(r[3])[-4:] for r in rows}, year_hint(Path(path).name))
    except Exception:
        return None


def find_templates(folder, out_root):
    found = []
    for p in sorted(Path(folder).glob("*.xlsx")):
        if (p.name.startswith("~$") or "Learners Performance" not in p.name
                or p.stem.endswith("- filled") or out_root in p.parents):
            continue
        t = peek_template(p)
        if t:
            found.append(t)
    return found


def best_job(tpl, jobs):
    """Which standard/year a downloaded template belongs to, and how that was decided."""
    cands = [j for j in jobs if j.standard == tpl.standard]
    if not cands:
        return None, ""
    if tpl.hint:
        hit = [j for j in cands if j.code == tpl.hint]
        return (hit[0], f"year {tpl.hint} in the file name") if hit else (None, "")
    if len(cands) == 1:
        return cands[0], "only year with this standard in the test data"

    def score(j):
        ids = set(j.learners)
        return (len(ids & tpl.ids) / len(ids), len(ids & tpl.ids))
    best = max(cands, key=score)
    return best, f"{len(set(best.learners) & tpl.ids)} of its {len(best.learners)} learners are in the file"


@dataclass
class Layout:
    score: dict        # (term, SEMIS subject) -> column
    level: dict        # (term, SEMIS subject) -> column
    remark: dict       # term -> column
    divisor: dict      # term -> number the subject total is divided by (x100)


def read_layout(ws):
    row1 = {c.column: c.value for c in ws[1] if c.value}
    starts = sorted(row1)
    lay = Layout({}, {}, {}, {})
    for i, col in enumerate(starts):
        m = re.fullmatch(r"Term\s*(\d)", str(row1[col]).strip(), re.I)
        if not m:
            continue
        term = int(m.group(1))
        end = starts[i + 1] - 1 if i + 1 < len(starts) else ws.max_column
        for c in range(col, end + 1):
            h = str(ws.cell(2, c).value or "").strip()
            if h.lower().startswith("subject grade"):
                lay.level[(term, h.split("-", 1)[1].strip())] = c
            elif h.lower() == "term remarks":
                lay.remark[term] = c
            elif h:
                lay.score[(term, h)] = c
    for term, c in lay.remark.items():
        f = str(ws.cell(4, c).value or "")
        m = re.search(r"/(\d+)\*100", f)
        n = sum(1 for (t, _) in lay.score if t == term)
        lay.divisor[term] = int(m.group(1)) if m else 100 * n
    return lay


def pick_score(rec, semis_subject, subject_map):
    for cand in subject_map.get(semis_subject, [semis_subject]):
        for subject, score in rec.scores.items():
            if norm(subject) == norm(cand):
                return subject, score
    return None, None


def fill_template(tpl, job, records, subject_map, out_dir, scores_rows, remark_rows, issues):
    wb = openpyxl.load_workbook(tpl.path)
    ws = data_sheet(wb)
    lay = read_layout(ws)
    if not lay.score:
        issues.append(("", "", job.label, "", f"{tpl.path.name}: no term/subject columns found; skipped"))
        return None, Counter()
    by_key = defaultdict(list)
    for r in records:
        by_key[(r.id, r.term)].append(r)
    stats, unmapped, blank_subject = Counter(), Counter(), Counter()
    seen_ids = set()
    for row in range(4, ws.max_row + 1):
        lin = ws.cell(row, 4).value
        if not lin:
            continue
        lid = str(lin)[-4:]
        for term in sorted({t for (t, _) in lay.score}):
            recs = by_key.get((lid, term))
            if not recs:
                continue
            if len(recs) > 1:
                issues.append((lid, recs[-1].name, job.label, term,
                               f"{len(recs)} mark rows for this learner and term; the last one was used"))
            rec = recs[-1]
            seen_ids.add(lid)
            used, entered = set(), []
            for (t, subject), col in sorted(lay.score.items(), key=lambda kv: kv[1]):
                if t != term:
                    continue
                source, score = pick_score(rec, subject, subject_map)
                if score is None:
                    blank_subject[subject] += 1
                    continue
                ws.cell(row, col).value = score
                used.add(source)
                entered.append(score)
                stats["scores"] += 1
                scores_rows.append((f"Tester {rec.tester}", lid, rec.name, rec.label, term, subject,
                                    source, score, level_of(score)))
            for subject in rec.scores:
                if subject not in used:
                    unmapped[subject] += 1
            for p in rec.problems:
                issues.append((lid, rec.name, job.label, term, p))
            if entered:
                avg = sum(entered) / lay.divisor[term] * 100
                exp = remark_of(avg)
                theirs = REMARK_EQUIV.get(str(rec.remark).strip().lower(), "?")
                remark_rows.append((f"Tester {rec.tester}", lid, rec.name, rec.label, term, len(entered),
                                    sum(entered), round(avg, 1), exp, rec.average, rec.remark,
                                    "Same" if exp == theirs else "Differs"))
                stats["learner_terms"] += 1
    for subject, n in sorted(unmapped.items()):
        issues.append(("", "", job.label, "", f"Test-data subject '{subject}' has no column in {tpl.path.name}: "
                       f"not entered ({n} learner-term(s))"))
    for subject, n in sorted(blank_subject.items()):
        issues.append(("", "", job.label, "", f"Template subject '{subject}' has no score in the test data for "
                       f"{n} learner-term(s): left blank"))
    for lid in sorted(tpl.ids - seen_ids):
        issues.append((lid, "", job.label, "", "Learner is in the template but has no marks for this standard/year "
                       "in the selected tester data (left blank; may belong to another tester)"))
    wb.calculation.fullCalcOnLoad = True    # grade levels and remarks are formulas: recalculate on open
    out = out_dir / f"{tpl.path.stem} - filled.xlsx"
    wb.save(out)
    stats["learners"] = len(seen_ids)
    return out, stats


# --------------------------------------------------------------------------- report
def write_report(job, summary, scores_rows, remark_rows, issues, out_dir):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 120
    for i, line in enumerate(summary, 1):
        ws.cell(i, 1, line)
    ws["A1"].font = Font(bold=True, size=14)
    sheet(wb, "Term remarks",
          ["Tester", "ID", "Learner", "Academic year", "Term", "Subjects entered", "Total",
           "Average (SEMIS basis)", "SEMIS remark should be", "Test data average", "Test data remark",
           "Remark vs test data"], sorted(remark_rows, key=lambda r: (r[0], r[1], r[4])))
    sheet(wb, "Scores entered",
          ["Tester", "ID", "Learner", "Academic year", "Term", "SEMIS subject", "Test-data subject",
           "Score", "SEMIS grade level should be"], sorted(scores_rows, key=lambda r: (r[0], r[1], r[4], r[5])))
    sheet(wb, "Issues and notes", ["ID", "Learner", "Academic year", "Term", "Note"],
          sorted(set(issues), key=lambda x: tuple(str(v) for v in x)))
    name = f"Performance report - {job.standard} of {job.code}.xlsx"
    wb.save(out_dir / name)
    return out_dir / name


# --------------------------------------------------------------------------- console
def show_job(job, cal):
    schools = Counter(r.school for r in {(r.id, r.school): r for r in job.records}.values())
    by_tester = Counter(t for t, _ in {(r.tester, r.id) for r in job.records})
    print()
    print("=" * 72)
    print(f" Download performance for {job.standard} of {job.code}  (academic year {job.label})")
    print(f"   Starting date : {dmy(cal.opens)}   (Term 1 opens)")
    print(f"   Closing date  : {dmy(cal.closes)}   (Term 3 closes)")
    if (cal.year_start, cal.year_end) != (cal.opens, cal.closes):
        print(f"   Calendar academic-year range is {dmy(cal.year_start)} to {dmy(cal.year_end)}")
    for t, (s, e) in sorted(cal.terms.items()):
        print(f"   Term {t}        : {dmy(s)} to {dmy(e)}")
    print(f"   Learners      : {len(job.learners)} ("
          + ", ".join(f"{s} {n}" for s, n in sorted(schools.items())) + ")")
    if len(by_tester) > 1:
        print("   Per tester    : " + ", ".join(f"Tester {t}: {n}" for t, n in sorted(by_tester.items())))
    print(f"   Tip           : save it with the year in the name, e.g. "
          f"'SEMIS - Learners Performance - {job.standard} - A - {job.code}.xlsx'")
    print("=" * 72)


def process_job(job, all_jobs, cal, args, subject_map, out_dir, out_root):
    def mine():
        out = []
        for t in find_templates(args.templates, out_root):
            j, how = best_job(t, all_jobs)
            if j is job:
                out.append((t, how))
        return out

    templates = mine()
    if not templates and not args.yes:
        print(f"\n No {job.standard} / {job.code} template found in {args.templates}")
        if ask(" Download it, save it there, then press Enter (or type s to skip): ", "") != "s":
            templates = mine()
    if not templates:
        print(" Skipped: no template available.")
        return
    scores_rows, remark_rows, issues = [], [], []
    summary = [f"Performance report: {job.standard} of {job.code} ({job.label})",
               f"Testers: {', '.join(f'Tester {t}' for t in sorted({r.tester for r in job.records}))}", "",
               f"Calendar: opens {dmy(cal.opens)}, closes {dmy(cal.closes)}."]
    covered = set()
    for tpl, how in templates:
        out, st = fill_template(tpl, job, job.records, subject_map, out_dir, scores_rows, remark_rows, issues)
        if out is None:
            continue
        covered |= tpl.ids
        summary += ["", f"Template: {tpl.path.name}  ->  {out.name}",
                    f"  Matched to {job.standard} of {job.code}: {how}",
                    f"  Learners filled : {st['learners']}",
                    f"  Scores entered  : {st['scores']}"]
        print(f" Filled {tpl.path.name} ({how}): {st['learners']} learners, {st['scores']} scores -> {out.name}")
    for lid in sorted(set(job.learners) - covered):
        issues.append((lid, "", job.label, "", f"Learner has {job.standard} {job.label} marks in the test data "
                       "but is not in any template found"))
    differs = sum(1 for r in remark_rows if r[-1] == "Differs")
    summary += ["", f"Learner-terms with a term remark: {len(remark_rows)}  (SEMIS remark differs from the test-data "
                f"remark in {differs}: the test data uses 5 remark bands, SEMIS 4)",
                f"Issues / notes: {len(set(issues))}  (see the 'Issues and notes' sheet)"]
    rep = write_report(job, summary, scores_rows, remark_rows, issues, out_dir)
    print(f" Notes: {len(set(issues))}   Report: {rep.name}")
    print(" Reminder: open the filled file in Excel and save once so the grade level / remark formulas "
          "store their results before uploading.")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    add_common_args(ap)
    ap.add_argument("--subject-map", help="JSON file {SEMIS subject: [test-data subjects]} overriding the defaults")
    args = ap.parse_args()
    calendars, available, picked, out_root, out_dir = prepare(args)
    subject_map = dict(SUBJECT_MAP)
    if args.subject_map:
        subject_map.update(json.loads(Path(args.subject_map).read_text(encoding="utf-8")))
    records = [r for n in picked for r in load_records(available[n], n)]
    all_jobs = build_jobs(records)
    jobs = [j for j in all_jobs if j.std >= args.start_standard]
    missing = sorted({j.code for j in jobs} - set(calendars))
    if missing:
        sys.exit(f"Calendar has no academic year for code(s): {', '.join(missing)}")

    for i, job in enumerate(jobs):
        show_job(job, calendars[job.code])
        process_job(job, all_jobs, calendars[job.code], args, subject_map, out_dir, out_root)
        if i + 1 < len(jobs):
            nxt = jobs[i + 1]
            if not args.yes and ask(f"\n Proceed with {nxt.standard} of {nxt.code}? [Y/n]: ") not in ("y", "yes"):
                print(" Stopped.")
                return
    print("\n All standards done. Output folder:", out_dir)


if __name__ == "__main__":
    main()
