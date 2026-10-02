#!/usr/bin/env python3
"""Guided bulk-attendance filler for the SEMIS lifecycle test.

Run it from the folder that holds your data (tester workbooks, calendar, templates).

For each standard (and academic year) in the tester's test data the script:
  1. tells you what to download from SEMIS (standard, year, start and closing date
     taken from "calendar script.json"),
  2. fills the downloaded "SEMIS - Learners Attendance - ..." template with
     `present` / `absent` using the Attendance sheet of the tester workbook,
  3. writes a change report for every date that had to be adjusted to fit the calendar,
  4. asks whether to continue with the next standard.

Usage:
    python fill_attendance.py                    # interactive: asks which tester(s), then guides each standard
    python fill_attendance.py --tester 2         # Tester 2 only (SEMIS_Test_Tester_2.xlsx)
    python fill_attendance.py --tester all       # all testers into the same templates (one upload per class)
    python fill_attendance.py --templates "C:\\Users\\me\\Downloads"
    python fill_attendance.py --report-only      # no templates: just produce the change reports
    python fill_attendance.py --yes              # never ask, skip standards whose template is missing

Rules applied (see the report for what changed):
  * The calendar decides term start/end, weekends and public holidays.
  * Template cells that are blank are school days; "Non School Day" cells are never touched.
  * A listed absent date that is not a school day (weekend, holiday, outside the term or
    outside the period the learner was at the school) is moved to the nearest free
    school day of the same term. The number of absences per learner is preserved.
  * "Late" is not a SEMIS option, so late days are recorded as `present`.
  * The tester workbook itself is never modified.
"""
import argparse
import datetime as dt
import json
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

WORK = Path.cwd()   # everything is looked up in the folder you run the script from
CALENDAR_NAMES = ("calendar.json", "calendar script.json")
PRESENT, ABSENT, NON_SCHOOL = "present", "absent", "Non School Day"
MONTHS = {m: i for i, m in enumerate(
    ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"], 1)}
WEEKDAYS = ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday"]
DATE_HEADER = re.compile(r"^\d{4}-\d{2}-\d{2}$")
SKIP_SHEETS = {"Validation", "Metadata"}


def dmy(d):
    return d.strftime("%d-%m-%Y")


def iso(s):
    return dt.date.fromisoformat(s)


# --------------------------------------------------------------------------- calendar
@dataclass
class Calendar:
    code: str
    label: str
    year_start: dt.date
    year_end: dt.date
    terms: dict          # {1: (start, end), ...}
    holidays: dict       # {date: event}
    weekdays: set        # {0..6}

    @property
    def opens(self):
        return min(s for s, _ in self.terms.values())

    @property
    def closes(self):
        return max(e for _, e in self.terms.values())

    def term_of(self, d):
        return next((n for n, (s, e) in self.terms.items() if s <= d <= e), None)

    def is_school_day(self, d):
        return (self.term_of(d) is not None and d.weekday() in self.weekdays
                and d not in self.holidays)

    def school_days(self, start=None, end=None):
        d, end = start or self.opens, end or self.closes
        while d <= end:
            if self.is_school_day(d):
                yield d
            d += dt.timedelta(days=1)


def load_calendar(path):
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    out = {}
    for e in data["schoolCalendar"]:
        ay = e["academicYear"]
        terms = {int(p["key"].replace("term", "")): (iso(p["startDate"]), iso(p["endDate"]))
                 for p in e["classPeriods"]}
        out[ay["code"]] = Calendar(
            code=ay["code"], label=ay["label"], year_start=iso(ay["startDate"]),
            year_end=iso(ay["endDate"]), terms=terms,
            holidays={iso(h["date"]): h["event"] for h in e["holidays"]},
            weekdays={i for i, n in enumerate(WEEKDAYS) if e["weekDays"].get(n)})
    return out


# --------------------------------------------------------------------------- test data
@dataclass
class Record:
    tester: int
    row: int
    id: str
    name: str
    label: str
    term: int
    standard: str
    school: str
    start: dt.date
    end: dt.date
    days: int
    late: int
    absent: int
    absent_on: list = field(default_factory=list)   # parsed dates
    late_on: list = field(default_factory=list)
    problems: list = field(default_factory=list)


def parse_days(text, label, problems, what):
    """'04 Oct, 09 Oct' -> dates; the year comes from the academic year label."""
    y1, y2 = int(label[:4]), int(label[5:])
    out = []
    for part in (text or "").split(","):
        part = part.strip()
        if not part:
            continue
        try:
            day, mon = part.split()
            month = MONTHS[mon[:3].title()]
            out.append(dt.date(y1 if month >= 8 else y2, month, int(day)))
        except (ValueError, KeyError):
            problems.append(f"Cannot read {what} date '{part}'")
    return out


def load_records(path, tester):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["Attendance"]
    out = []
    for n, r in enumerate(ws.iter_rows(min_row=5, values_only=True), start=5):
        if not r or not r[0]:
            continue
        rec = Record(tester=tester, row=n, id=str(r[0]), name=r[1], label=r[2], term=int(r[3]), standard=r[4],
                     school=r[5], start=r[6].date(), end=r[7].date(), days=int(r[8] or 0),
                     late=int(r[10] or 0), absent=int(r[11] or 0))
        rec.absent_on = parse_days(r[12], rec.label, rec.problems, "absent")
        rec.late_on = parse_days(r[13], rec.label, rec.problems, "late")
        out.append(rec)
    classes = {}
    if "Enrolment plan" in wb.sheetnames:
        for r in wb["Enrolment plan"].iter_rows(min_row=5, values_only=True):
            if r and r[0]:
                classes[(tester, str(r[0]), r[2])] = r[5]
    wb.close()
    return out, classes


def discover_testers(folder):
    """{1: Path('SEMIS_Test_Tester_1.xlsx'), ...} for every tester workbook in the folder."""
    found = {}
    for p in Path(folder).glob("SEMIS_Test_Tester_*.xlsx"):
        m = re.fullmatch(r"SEMIS_Test_Tester_(\d+)\.xlsx", p.name, re.I)
        if m:
            found[int(m.group(1))] = p
    return dict(sorted(found.items()))


def choose_testers(available, wanted):
    """wanted: None (ask), 'all' or '1,3,4'."""
    if not available:
        sys.exit("No tester workbooks (SEMIS_Test_Tester_<n>.xlsx) found. Use --data-dir.")
    if wanted is None:
        if len(available) == 1:
            return list(available)
        print("\n Tester workbooks found: " + ", ".join(f"Tester {n}" for n in available))
        wanted = ask(" Which tester(s)? e.g. 1  or  1,3  or  all [all]: ", "all")
    if wanted.strip().lower() == "all":
        return list(available)
    try:
        picked = sorted({int(x) for x in re.split(r"[,\s]+", wanted.strip()) if x})
    except ValueError:
        sys.exit(f"Cannot read tester selection '{wanted}'")
    bad = [n for n in picked if n not in available]
    if bad or not picked:
        sys.exit(f"No workbook for tester(s) {bad or wanted}; available: {list(available)}")
    return picked


def code_of(label):
    return label[-4:]


def std_num(standard):
    return int(re.search(r"\d+", standard).group())


@dataclass
class Job:
    std: int
    standard: str
    label: str
    code: str
    records: list

    @property
    def learners(self):
        return sorted({r.id for r in self.records})


def build_jobs(records):
    groups = defaultdict(list)
    for r in records:
        groups[(std_num(r.standard), code_of(r.label), r.standard, r.label)].append(r)
    return [Job(k[0], k[2], k[3], k[1], v) for k, v in sorted(groups.items())]


# --------------------------------------------------------------------------- planning
@dataclass
class Segment:
    rec: Record
    start: dt.date
    end: dt.date
    absent: set


def term_reference(records):
    """Most common (from, to) per (tester, year, term): what a learner present all term has."""
    c = defaultdict(Counter)
    for r in records:
        c[(r.tester, r.label, r.term)][(r.start, r.end)] += 1
    return {k: v.most_common(1)[0][0] for k, v in c.items()}


def why_moved(d, cal, seen, start, end, term_range):
    if d in seen:
        return "Duplicate date in test data"
    if d.weekday() not in cal.weekdays:
        return f"{d.strftime('%A')} (weekend)"
    if d in cal.holidays:
        return f"Public holiday: {cal.holidays[d]}"
    if not (term_range[0] <= d <= term_range[1]):
        return f"Outside calendar term ({dmy(term_range[0])} to {dmy(term_range[1])})"
    if not (start <= d <= end):
        return f"Outside the period learner was at the school ({dmy(start)} to {dmy(end)})"
    return "Not a school day in the template"


def plan_records(records, cal, school_days, ref, adjustments, issues):
    """Turn test records into Segments (who is present/absent on which day)."""
    segments = defaultdict(list)
    for rec in records:
        t_start, t_end = cal.terms[rec.term]
        r_from, r_to = ref[(rec.tester, rec.label, rec.term)]
        # Full-term learners follow the calendar; transfers/dropouts keep their own dates.
        start = max(rec.start if rec.start > r_from else t_start, t_start)
        end = min(rec.end if rec.end < r_to else t_end, t_end)
        window = sorted(d for d in school_days if start <= d <= end)
        wset, taken, pending = set(window), set(), []
        for d in sorted(rec.absent_on):
            (taken.add(d) if d in wset and d not in taken else pending.append(d))
        seen = set(taken)
        for d in pending:
            reason = why_moved(d, cal, seen, start, end, (t_start, t_end))
            free = [x for x in window if x not in taken]
            if not free:
                issues.append((rec.id, rec.name, rec.label, rec.term,
                               f"No free school day to move absence of {dmy(d)} ({reason})"))
                continue
            late = set(rec.late_on)
            new = min(free, key=lambda x: (x in late, abs((x - d).days), x))
            taken.add(new)
            adjustments.append(dict(tester=rec.tester, id=rec.id, name=rec.name, label=rec.label, term=rec.term,
                                    school=rec.school, old=d, new=new, reason=reason))
        for p in rec.problems:
            issues.append((rec.id, rec.name, rec.label, rec.term, p))
        if rec.absent != len(rec.absent_on):
            issues.append((rec.id, rec.name, rec.label, rec.term,
                           f"Absent count is {rec.absent} but {len(rec.absent_on)} dates are listed "
                           "(dates were used)"))
        segments[rec.id].append(Segment(rec, start, end, taken))
    return segments


# --------------------------------------------------------------------------- templates
@dataclass
class Template:
    path: Path
    code: str
    standard: str
    stream: str
    school: str


def month_sheets(wb):
    return [ws for ws in wb.worksheets if ws.title not in SKIP_SHEETS]


def peek_template(path):
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheets = [n for n in wb.sheetnames if n not in SKIP_SHEETS]
        row = next(wb[sheets[0]].iter_rows(min_row=4, max_row=4, values_only=True))
        wb.close()
        return Template(Path(path), str(row[3]), str(row[4]), str(row[5]), str(row[1]))
    except Exception:
        return None


def find_templates(folder, out_dir):
    found = []
    for p in sorted(Path(folder).glob("*.xlsx")):
        if (p.name.startswith("~$") or "Learners Attendance" not in p.name
                or p.stem.endswith("- filled") or out_dir in p.parents):
            continue
        t = peek_template(p)
        if t:
            found.append(t)
    return found


def template_days(wb):
    """Dates whose cell is empty (= school day) in the first learner row, plus date->column maps."""
    days, cols = set(), {}
    for ws in month_sheets(wb):
        cols[ws.title] = {c.column: iso(c.value) for c in ws[2]
                          if isinstance(c.value, str) and DATE_HEADER.match(c.value)}
        for col, d in cols[ws.title].items():
            if ws.cell(4, col).value in (None, ""):
                days.add(d)
    return days, cols


def fill_template(tpl, job, cal, records, classes, out_dir, adjustments, issues):
    wb = openpyxl.load_workbook(tpl.path)
    days, cols = template_days(wb)
    cal_days = set(cal.school_days())
    diff = sorted(days ^ cal_days)
    if diff:
        issues.append(("", "", job.label, "", f"Template school days differ from the calendar on "
                       f"{len(diff)} date(s), e.g. {', '.join(dmy(d) for d in diff[:5])}. "
                       "The template was followed."))
    first = month_sheets(wb)[0]
    in_tpl = {str(first.cell(r, 7).value)[-4:] for r in range(4, first.max_row + 1)
              if first.cell(r, 7).value}
    mine = [r for r in records if r.id in in_tpl]
    owners = defaultdict(set)
    for r in mine:
        owners[r.id].add(r.tester)
    for lid, t in owners.items():
        if len(t) > 1:
            issues.append((lid, "", job.label, "", "Learner ID appears in more than one tester workbook "
                           f"({', '.join(f'Tester {x}' for x in sorted(t))}); check for a clash"))
    ref = term_reference(records)
    segments = plan_records(mine, cal, days, ref, adjustments, issues)

    for i in sorted(in_tpl - {r.id for r in records}):
        issues.append((i, "", job.label, "", f"Learner is in the template but has no {job.standard} "
                       f"{job.label} attendance in the selected tester data (left blank; may belong to another tester)"))
    stats = Counter()
    for ws in month_sheets(wb):
        for r in range(4, ws.max_row + 1):
            lin = ws.cell(r, 7).value
            segs = segments.get(str(lin)[-4:]) if lin else None
            if not segs:
                continue
            for col, d in cols[ws.title].items():
                cell = ws.cell(r, col)
                if cell.value == NON_SCHOOL:
                    continue
                seg = next((s for s in segs if s.start <= d <= s.end), None)
                if seg is None:
                    continue
                cell.value = ABSENT if d in seg.absent else PRESENT
                stats[cell.value] += 1
    other_school, other_class = defaultdict(set), {}
    for lid, segs in segments.items():
        for s in segs:
            if tpl.school.lower().find(s.rec.school.lower()) < 0:
                other_school[(lid, s.rec.name, s.rec.school)].add(s.rec.term)
            cls = classes.get((s.rec.tester, lid, s.rec.label))
            if cls and cls != tpl.stream:
                other_class[lid] = cls
    for (lid, name, school), terms in other_school.items():
        issues.append((lid, name, job.label, ", ".join(map(str, sorted(terms))),
                       f"Test data says {school} but the template is '{tpl.school}'; "
                       "filled anyway (check the learner's enrolment/transfer in SEMIS)"))
    if other_class:
        issues.append(("", "", job.label, "", f"Template stream is {tpl.stream}; the test data class differs for "
                       f"{len(other_class)} learner(s): {', '.join(sorted(other_class))}. Matched by learner ID."))
    late_days = sum(len([d for d in s.rec.late_on if s.start <= d <= s.end])
                    for ss in segments.values() for s in ss)
    stats["late_as_present"] = late_days
    stats["learners"] = len(segments)
    out = out_dir / f"{tpl.path.stem} - filled.xlsx"
    wb.save(out)
    return out, stats


# --------------------------------------------------------------------------- report
HEAD = PatternFill("solid", fgColor="1F4E78")


def sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(list(r))
    for i, h in enumerate(header, 1):
        w = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r)] + [8])
        ws.column_dimensions[get_column_letter(i)].width = min(w + 2, (widths or {}).get(i, 70))
    ws.freeze_panes = "A2"
    return ws


def write_report(job, cal, adjustments, issues, summary_lines, all_records, out_dir):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 120
    for i, line in enumerate(summary_lines, 1):
        ws.cell(i, 1, line)
    ws["A1"].font = Font(bold=True, size=14)

    sheet(wb, "Date changes",
          ["Tester", "ID", "Learner", "Academic year", "Term", "School", "Original absent date", "Day",
           "New absent date", "Day", "Days moved", "Reason"],
          [(f"Tester {a['tester']}", a["id"], a["name"], a["label"], a["term"], a["school"], dmy(a["old"]),
            a["old"].strftime("%a"), dmy(a["new"]), a["new"].strftime("%a"),
            (a["new"] - a["old"]).days, a["reason"])
           for a in sorted(adjustments, key=lambda a: (a["tester"], a["id"], a["term"], a["old"]))])

    ref = term_reference(all_records)
    per = []
    for (tester, label, term), (f, t) in sorted(ref.items()):
        if label != job.label:
            continue
        s, e = cal.terms[term]
        reported = [r.days for r in all_records if r.tester == tester and r.label == label
                    and r.term == term and (r.start, r.end) == (f, t)][0]
        n = len(list(cal.school_days(s, e)))
        per.append((f"Tester {tester}", label, term, dmy(f), dmy(t), reported, dmy(s), dmy(e), n,
                    "OK" if (f, t, reported) == (s, e, n) else "Differs - calendar used"))
    sheet(wb, "Term dates check",
          ["Tester", "Academic year", "Term", "Test data from", "Test data to", "Test data school days",
           "Calendar from", "Calendar to", "Calendar school days", "Status"], per)

    sheet(wb, "Issues and notes", ["ID", "Learner", "Academic year", "Term", "Note"],
          sorted(set(issues), key=lambda x: tuple(str(v) for v in x)))
    name = f"Attendance report - {job.standard} of {job.code}.xlsx"
    wb.save(out_dir / name)
    return out_dir / name


# --------------------------------------------------------------------------- console
def ask(prompt, default="y"):
    try:
        ans = input(prompt).strip().lower()
    except EOFError:
        return "n"
    return ans or default


def show_job(job, cal):
    schools = Counter(r.school for r in {(r.id, r.school): r for r in job.records}.values())
    print()
    print("=" * 72)
    print(f" Download attendance for {job.standard} of {job.code}  (academic year {job.label})")
    print(f"   Starting date : {dmy(cal.opens)}   (Term 1 opens)")
    print(f"   Closing date  : {dmy(cal.closes)}   (Term 3 closes)")
    if (cal.year_start, cal.year_end) != (cal.opens, cal.closes):
        print(f"   Calendar academic-year range is {dmy(cal.year_start)} to {dmy(cal.year_end)}")
    print(f"   Learners      : {len(job.learners)} ("
          + ", ".join(f"{s} {n}" for s, n in sorted(schools.items())) + ")")
    by_tester = Counter(t for t, _ in {(r.tester, r.id) for r in job.records})
    if len(by_tester) > 1:
        print("   Per tester    : " + ", ".join(f"Tester {t}: {n}" for t, n in sorted(by_tester.items())))
    print("=" * 72)


def process_job(job, cal, args, classes, all_records, out_dir, out_root):
    adjustments, issues, lines = [], [], []
    lines += [f"Attendance report: {job.standard} of {job.code} ({job.label})",
              f"Testers: {', '.join(f'Tester {t}' for t in sorted({r.tester for r in job.records}))}", "",
              f"Calendar: opens {dmy(cal.opens)}, closes {dmy(cal.closes)}; "
              f"public holidays and weekends are non-school days."]
    templates = []
    if not args.report_only:
        templates = [t for t in find_templates(args.templates, out_root)
                     if t.code == job.code and t.standard == job.standard]
        if not templates and not args.yes:
            print(f"\n No {job.standard} / {job.code} template found in {args.templates}")
            ans = ask(" Download it, save it there, then press Enter (or type s to skip): ", "")
            if ans != "s":
                templates = [t for t in find_templates(args.templates, out_root)
                             if t.code == job.code and t.standard == job.standard]
        if not templates:
            print(" Skipped: no template available.")
            return
    if args.report_only:
        days = set(cal.school_days())
        plan_records(job.records, cal, days, term_reference(all_records), adjustments, issues)
        lines.append("Mode: report only (school days taken from the calendar, no template filled).")
    for tpl in templates:
        out, st = fill_template(tpl, job, cal, job.records, classes, out_dir, adjustments, issues)
        lines += ["", f"Template: {tpl.path.name}  ->  {out.name}",
                  f"  Learners filled : {st['learners']}",
                  f"  Cells present   : {st[PRESENT]}",
                  f"  Cells absent    : {st[ABSENT]}",
                  f"  Late days recorded as present (SEMIS has no 'late'): {st['late_as_present']}"]
        print(f" Filled {tpl.path.name}: {st['learners']} learners, {st[PRESENT]} present, "
              f"{st[ABSENT]} absent -> {out.name}")
    lines += ["", f"Absent dates moved to fit the calendar: {len(adjustments)}",
              f"Issues / notes: {len(set(issues))}  (see the 'Issues and notes' sheet)"]
    rep = write_report(job, cal, adjustments, issues, lines, all_records, out_dir)
    print(f" Dates adjusted: {len(adjustments)}   Notes: {len(set(issues))}   Report: {rep.name}")


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--data-dir", default=WORK, help="folder with SEMIS_Test_Tester_<n>.xlsx files")
    ap.add_argument("--tester", help="tester number(s): 1, 1,3 or all (asked if omitted)")
    ap.add_argument("--calendar", help="school calendar JSON (default: calendar.json or "
                    "'calendar script.json' in the current folder)")
    ap.add_argument("--templates", default=WORK, help="folder holding the downloaded templates")
    ap.add_argument("--out", default=WORK / "output", help="folder for filled templates and reports")
    ap.add_argument("--start-standard", type=int, default=1)
    ap.add_argument("--yes", action="store_true", help="do not ask; skip missing templates")
    ap.add_argument("--report-only", action="store_true", help="only build change reports")
    args = ap.parse_args()
    args.templates, out_root = Path(args.templates), Path(args.out)

    cal_path = Path(args.calendar) if args.calendar else next(
        (WORK / n for n in CALENDAR_NAMES if (WORK / n).exists()), None)
    if not cal_path or not cal_path.exists():
        sys.exit("Calendar not found. Put calendar.json in this folder or pass --calendar <file>.")
    calendars = load_calendar(cal_path)
    available = discover_testers(args.data_dir)
    picked = choose_testers(available, args.tester)
    records, classes = [], {}
    for n in picked:
        recs, cls = load_records(available[n], n)
        records += recs
        classes.update(cls)
    if len(picked) == 1:
        label = f"Tester {picked[0]}"
    elif len(picked) == len(available):
        label = "All testers"
    else:
        label = "Testers " + ",".join(map(str, picked))
    out_dir = out_root / label
    out_dir.mkdir(parents=True, exist_ok=True)
    print(f"\n Using: {', '.join(f'Tester {n}' for n in picked)}   Output: {out_dir}")
    jobs = [j for j in build_jobs(records) if j.std >= args.start_standard]
    missing = sorted({j.code for j in jobs} - set(calendars))
    if missing:
        sys.exit(f"Calendar has no academic year for code(s): {', '.join(missing)}")

    for i, job in enumerate(jobs):
        cal = calendars[job.code]
        show_job(job, cal)
        process_job(job, cal, args, classes, records, out_dir, out_root)
        if i + 1 < len(jobs):
            nxt = jobs[i + 1]
            if not args.yes and not args.report_only and ask(
                    f"\n Proceed with {nxt.standard} of {nxt.code}? [Y/n]: ") not in ("y", "yes"):
                print(" Stopped.")
                return
    print("\n All standards done. Output folder:", out_dir)


if __name__ == "__main__":
    main()
