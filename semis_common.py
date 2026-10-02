"""Helpers shared by the SEMIS test-data loaders (attendance/, performance/)."""
import datetime as dt
import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

WORK = Path.cwd()   # everything is looked up in the folder you run a script from
CALENDAR_NAMES = ("calendar.json", "calendar script.json")
WEEKDAYS = ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday"]
SKIP_SHEETS = {"Validation", "Metadata"}


def dmy(d):
    return d.strftime("%d-%m-%Y")


def iso(s):
    return dt.date.fromisoformat(s)


def code_of(label):
    """'2018-2019' -> '2019' (the academic year code SEMIS uses)."""
    return label[-4:]


def std_num(standard):
    return int(re.search(r"\d+", standard).group())


def ask(prompt, default="y"):
    try:
        ans = input(prompt).strip().lower()
    except EOFError:
        return "n"
    return ans or default


# --------------------------------------------------------------------------- calendar
@dataclass
class Calendar:
    code: str
    label: str
    year_start: dt.date
    year_end: dt.date
    terms: dict          # {1: (start, end), ...}
    holidays: dict       # {date: event}
    weekdays: set        # {0..6}

    @property
    def opens(self):
        return min(s for s, _ in self.terms.values())

    @property
    def closes(self):
        return max(e for _, e in self.terms.values())

    def term_of(self, d):
        return next((n for n, (s, e) in self.terms.items() if s <= d <= e), None)

    def is_school_day(self, d):
        return (self.term_of(d) is not None and d.weekday() in self.weekdays
                and d not in self.holidays)

    def school_days(self, start=None, end=None):
        d, end = start or self.opens, end or self.closes
        while d <= end:
            if self.is_school_day(d):
                yield d
            d += dt.timedelta(days=1)


def load_calendar(path):
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    out = {}
    for e in data["schoolCalendar"]:
        ay = e["academicYear"]
        terms = {int(p["key"].replace("term", "")): (iso(p["startDate"]), iso(p["endDate"]))
                 for p in e["classPeriods"]}
        out[ay["code"]] = Calendar(
            code=ay["code"], label=ay["label"], year_start=iso(ay["startDate"]),
            year_end=iso(ay["endDate"]), terms=terms,
            holidays={iso(h["date"]): h["event"] for h in e["holidays"]},
            weekdays={i for i, n in enumerate(WEEKDAYS) if e["weekDays"].get(n)})
    return out


def find_calendar(explicit, *folders):
    if explicit:
        return Path(explicit)
    for folder in folders:
        for name in CALENDAR_NAMES:
            if (Path(folder) / name).exists():
                return Path(folder) / name
    sys.exit("Calendar not found. Put calendar.json next to your data or pass --calendar <file>.")


# --------------------------------------------------------------------------- testers
def discover_testers(folder):
    """{1: Path('SEMIS_Test_Tester_1.xlsx'), ...} for every tester workbook in the folder."""
    found = {}
    for p in Path(folder).glob("SEMIS_Test_Tester_*.xlsx"):
        m = re.fullmatch(r"SEMIS_Test_Tester_(\d+)\.xlsx", p.name, re.I)
        if m:
            found[int(m.group(1))] = p
    return dict(sorted(found.items()))


def choose_testers(available, wanted):
    """wanted: None (ask), 'all' or '1,3,4'."""
    if not available:
        sys.exit("No tester workbooks (SEMIS_Test_Tester_<n>.xlsx) found. Use --data-dir.")
    if wanted is None:
        if len(available) == 1:
            return list(available)
        print("\n Tester workbooks found: " + ", ".join(f"Tester {n}" for n in available))
        wanted = ask(" Which tester(s)? e.g. 1  or  1,3  or  all [all]: ", "all")
    if wanted.strip().lower() == "all":
        return list(available)
    try:
        picked = sorted({int(x) for x in re.split(r"[,\s]+", wanted.strip()) if x})
    except ValueError:
        sys.exit(f"Cannot read tester selection '{wanted}'")
    bad = [n for n in picked if n not in available]
    if bad or not picked:
        sys.exit(f"No workbook for tester(s) {bad or wanted}; available: {list(available)}")
    return picked


def add_common_args(ap):
    ap.add_argument("--data-dir", default=WORK, help="folder with SEMIS_Test_Tester_<n>.xlsx files")
    ap.add_argument("--tester", help="tester number(s): 1, 1,3 or all (asked if omitted)")
    ap.add_argument("--calendar", help="school calendar JSON (default: calendar.json or "
                    "'calendar script.json' in --data-dir, then the current folder)")
    ap.add_argument("--templates", default=WORK, help="folder holding the downloaded templates")
    ap.add_argument("--out", default=WORK / "output", help="folder for filled templates and reports")
    ap.add_argument("--start-standard", type=int, default=1)
    ap.add_argument("--yes", action="store_true", help="do not ask; skip missing templates")


def prepare(args):
    """Resolve folders, calendar and tester selection.

    Returns (calendars, available, picked, out_root, out_dir); out_dir exists.
    """
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    args.templates, out_root = Path(args.templates), Path(args.out)
    calendars = load_calendar(find_calendar(args.calendar, args.data_dir, WORK))
    available = discover_testers(args.data_dir)
    picked = choose_testers(available, args.tester)
    if len(picked) == 1:
        label = f"Tester {picked[0]}"
    elif len(picked) == len(available):
        label = "All testers"
    else:
        label = "Testers " + ",".join(map(str, picked))
    out_dir = out_root / label
    out_dir.mkdir(parents=True, exist_ok=True)
    print(f"\n Using: {', '.join(f'Tester {n}' for n in picked)}   Output: {out_dir}")
    return calendars, available, picked, out_root, out_dir


# --------------------------------------------------------------------------- reports
HEAD = PatternFill("solid", fgColor="1F4E78")


def sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(list(r))
    for i, h in enumerate(header, 1):
        w = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r)] + [8])
        ws.column_dimensions[get_column_letter(i)].width = min(w + 2, (widths or {}).get(i, 70))
    ws.freeze_panes = "A2"
    return ws
